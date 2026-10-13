import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_left_border(cell, hex_color="0D9488", sz="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_callout(doc, text, title=None, hex_bg="F0FDF4", border_color="059669"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, hex_bg)
    set_cell_left_border(cell, border_color, sz="30")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(title + "\n")
        run_title.bold = True
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = RGBColor(15, 118, 110)
    
    run_txt = p.add_run(text)
    run_txt.font.name = "Segoe UI"
    run_txt.font.size = Pt(10)
    run_txt.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_run(run, font_name="Segoe UI", size_pt=10.5, color_rgb=(51, 65, 85), bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    run.bold = bold
    run.italic = italic

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    
    run = h.add_run(text)
    format_run(run, font_name="Segoe UI", size_pt=14, color_rgb=(15, 23, 42), bold=True)
    
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="0D9488"/></w:pBdr>')
    h._p.get_or_add_pPr().append(pBdr)
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    format_run(run, font_name="Segoe UI", size_pt=11.5, color_rgb=(13, 148, 136), bold=True)
    return h

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix + ": ")
        format_run(r1, font_name="Segoe UI", size_pt=10.5, color_rgb=(15, 23, 42), bold=True)
    r2 = p.add_run(text)
    format_run(r2, font_name="Segoe UI", size_pt=10.5, color_rgb=(51, 65, 85))
    return p

def add_sub_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet 2')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix + ": ")
        format_run(r1, font_name="Segoe UI", size_pt=10, color_rgb=(30, 41, 59), bold=True)
    r2 = p.add_run(text)
    format_run(r2, font_name="Segoe UI", size_pt=10, color_rgb=(71, 85, 105))
    return p

def create_brief_docx(filename):
    doc = docx.Document()
    
    # Page Setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        s.header.is_linked_to_previous = False
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("TÀI LIỆU TÓM TẮT DỰ ÁN | DRIVERENT")
        format_run(hrun, font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Hệ thống DriveRent · 8 Câu Hỏi Định Vị Bài Toán Trước Khi Làm")
        format_run(frun, font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))

    # Document Header Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("TÀI LIỆU TÓM TẮT DỰ ÁN")
    format_run(t_run, font_name="Segoe UI", size_pt=20, color_rgb=(15, 23, 42), bold=True)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    s_run = sub_p.add_run("HỆ THỐNG QUẢN LÝ DỊCH VỤ CHO THUÊ XE MÁY & Ô TÔ TỰ LÁI DRIVERENT\nChuẩn phương pháp: Hiểu đúng bài toán trước khi làm (8 câu hỏi cốt lõi)")
    format_run(s_run, font_name="Segoe UI", size_pt=11, color_rgb=(13, 148, 136), bold=True)

    # Info Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    col_widths = [Inches(2.0), Inches(4.7)]
    
    meta_data = [
        ("Tên dự án:", "DriveRent - Hệ thống Quản lý Cho thuê Xe máy & Ô tô Tự lái Thông minh"),
        ("Mã dự án & Phiên bản:", "DRV-RENT-2026 | Phiên bản 1.0.0"),
        ("Kho lưu trữ mã nguồn:", "https://github.com/DinhTuanHai-TBD/QuanLyDatXe_DriveRent.git"),
        ("Ngày cập nhật:", "13/10/2026 (Đồng bộ chuẩn 8 câu hỏi cốt lõi)")
    ]
    
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=70, bottom=70, left=120, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        format_run(r0, font_name="Segoe UI", size_pt=9.5, color_rgb=(71, 85, 105), bold=True)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        format_run(r1, font_name="Segoe UI", size_pt=9.5, color_rgb=(15, 23, 42))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Callout overview
    add_callout(
        doc,
        "Brief là bảng mô tả tóm tắt thông tin của một dự án công nghệ, đồng thời là căn cứ cốt lõi để tư vấn, định vị bài toán và trao đổi với khách hàng. Một bản brief chuẩn trả lời trọn vẹn tám câu hỏi phân bổ thành 2 nhóm: Nhóm định vị khách hàng và mục tiêu (Câu 1 đến Câu 4) và Nhóm giải pháp kỹ thuật, phạm vi và kết quả (Câu 5 đến Câu 8).",
        title="ĐỊNH NGHĨA & NGUYÊN TẮC THIẾT KẾ BẢN TÓM TẮT DỰ ÁN",
        hex_bg="F0FDFA",
        border_color="0D9488"
    )

    # ==========================
    # CÂU 1: DỰ ÁN NÀY LÀ GÌ?
    # ==========================
    add_heading_1(doc, "1. Dự án này là gì?")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("DriveRent là Website và Hệ thống phần mềm quản lý dịch vụ cho thuê xe máy và ô tô tự lái toàn diện. Dự án tích hợp đồng bộ hai phân hệ chuyên biệt trên cùng một nền tảng:")
    format_run(r, font_name="Segoe UI", size_pt=10.5)

    add_bullet(doc, "Cổng Khách hàng", 
               "Website trực tuyến hiện đại dành cho du khách tìm kiếm, tham khảo thông số, kiểm tra giá thuê theo giờ hoặc theo ngày và đặt cọc giữ chỗ trực tuyến. Ứng dụng ngôn ngữ thiết kế kính mờ trong suốt cao cấp, hiệu ứng xe tự động chạy vào bãi gara đèn LED sống động, thanh điều hướng chuyển đổi trạng thái mượt mà theo vị trí cuộn trang.")
    
    add_bullet(doc, "Hệ thống Quản trị và Điều hành Nội bộ", 
               "Bảng điều khiển dành riêng cho chủ nhà xe, nhân viên lễ tân và thợ kỹ thuật. Giám sát đội xe trực quan theo thời gian thực gồm Sẵn sàng, Đang thuê, Bảo dưỡng; số hóa quy trình lập hợp đồng trong 2 phút; quản lý biên bản bàn giao và thu hồi xe kèm ảnh chụp hiện trạng; tự động quyết toán phân định tiền cọc và quản lý lịch bảo trì, nhắc thay nhớt định kỳ.")

    # ==========================
    # CÂU 2: KHÁCH HÀNG LÀ AI?
    # ==========================
    add_heading_1(doc, "2. Khách hàng là ai?")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("Khách hàng trực tiếp của dự án là:")
    format_run(r, font_name="Segoe UI", size_pt=10.5)

    add_bullet(doc, "Cơ sở và Chuỗi cửa hàng kinh doanh cho thuê phương tiện tự lái",
               "Các nhà xe có quy mô vừa và nhỏ, sở hữu từ 15 đến hơn 80 đầu xe hỗn hợp cả xe máy tay ga, xe số và ô tô 4 đến 7 chỗ, tập trung tại các thành phố du lịch trọng điểm của Việt Nam như Đà Nẵng, Nha Trang, Phú Quốc, Đà Lạt, Quy Nhơn, Vũng Tàu, Hà Nội, TP.HCM,...")

    add_bullet(doc, "Đội ngũ nhân sự vận hành tại nhà xe",
               "Các bộ phận phối hợp nhịp nhàng trong quy trình:")
    add_sub_bullet(doc, "Chủ cửa hàng, Quản lý", "Cần giám sát tổng thể doanh thu, công nợ, dòng tiền cọc đang nắm giữ và hiệu suất khai thác đội xe mọi lúc mọi nơi.")
    add_sub_bullet(doc, "Nhân viên Lễ tân, Tiếp đón khách", "Cần thao tác tra cứu xe trống tức thời, lập hợp đồng thuê xe, bàn giao xe và thanh lý hợp đồng chính xác chỉ trong vài phút.")
    add_sub_bullet(doc, "Kỹ thuật viên, Thợ bảo dưỡng", "Cần theo dõi số km lăn bánh của từng xe để thay nhớt, kiểm tra má phanh, lốp xe đúng định kỳ.")

    # ==========================
    # CÂU 3: ĐỐI TƯỢNG KHÁCH HÀNG CỦA HỌ LÀ AI?
    # ==========================
    add_heading_1(doc, "3. Đối tượng khách hàng của họ là ai?")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("Đối tượng khách hàng mục tiêu của nhà xe gồm 4 nhóm chính:")
    format_run(r, font_name="Segoe UI", size_pt=10.5)

    add_bullet(doc, "Khách du lịch tự túc",
               "Cá nhân, cặp đôi hoặc nhóm bạn trẻ đến các điểm du lịch nghỉ dưỡng từ 2 đến 5 ngày, muốn thuê xe máy như Vision, AirBlade hoặc ô tô 4 chỗ như Vios để chủ động lịch trình khám phá. Ưa chuộng hình ảnh thật, bảng giá minh bạch và đặt cọc giữ xe nhanh qua mã QR chuyển khoản.")
    add_bullet(doc, "Khách du lịch gia đình",
               "Đoàn từ 4 đến 7 người, ưu tiên thuê ô tô 7 chỗ như Mitsubishi Xpander, Toyota Veloz đời mới, sạch sẽ, khoang hành lý rộng, điều hòa mát và độ an toàn kỹ thuật cao.")
    add_bullet(doc, "Chuyên gia và Khách công tác ngắn ngày",
               "Cần thuê ô tô tự lái theo giờ hoặc theo ngày để đi gặp đối tác, di chuyển từ sân bay về khách sạn với hóa đơn, hợp đồng chuẩn A4.")
    add_bullet(doc, "Người dân địa phương",
               "Thuê xe tự lái phục vụ nhu cầu đi lại dịp lễ Tết, cưới hỏi, về quê hoặc khi phương tiện cá nhân đang trong xưởng bảo dưỡng.")

    # ==========================
    # CÂU 4: MỤC TIÊU CỦA DỰ ÁN LÀ GÌ?
    # ==========================
    add_heading_1(doc, "4. Mục tiêu của dự án là gì?")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("Dự án DriveRent được triển khai nhằm đạt được 5 mục tiêu chiến lược sau:")
    format_run(r, font_name="Segoe UI", size_pt=10.5)

    add_bullet(doc, "1. Số hóa 100% quy trình vận hành", 
               "Thay thế hoàn toàn sổ ghi tay, file Excel rời rạc hoặc tin nhắn dễ trôi bằng hệ thống dữ liệu số tập trung.")
    add_bullet(doc, "2. Rút ngắn thời gian xử lý thủ tục", 
               "Giảm thời gian làm hợp đồng và bàn giao xe từ 15–20 phút xuống còn dưới 2–3 phút mỗi khách nhờ quy trình 4 bước và biên bản kiểm tra số hóa.")
    add_bullet(doc, "3. Minh bạch tài chính và Phân định tiền cọc", 
               "Cách ly hoàn toàn tiền đặt cọc khỏi doanh thu dịch vụ, tự động tính cước thuê thực tế và khấu trừ cọc minh bạch, chấm dứt tranh cãi lúc hoàn cọc.")
    add_bullet(doc, "4. Bảo vệ tài sản và Kéo dài tuổi thọ phương tiện", 
               "Số hóa lưu trữ ảnh căn cước công dân, bằng lái xe và ảnh hiện trạng xe lúc giao và nhận; tự động nhắc thay nhớt theo số km thực tế giúp kéo dài tuổi thọ động cơ và đảm bảo an toàn cho khách.")
    add_bullet(doc, "5. Đột phá trải nghiệm và Định vị thương hiệu", 
               "Cung cấp giao diện đặt xe phong cách kính mờ trong suốt sang trọng, hiệu ứng sinh động, nâng cao tỷ lệ chuyển đổi khách đặt xe trực tuyến.")

    # ==========================
    # CÂU 5: WEBSITE CẦN GIẢI QUYẾT VẤN ĐỀ GÌ?
    # ==========================
    add_heading_1(doc, "5. Website cần giải quyết vấn đề gì?")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("Website DriveRent được thiết kế để giải quyết triệt để 4 khó khăn và bất cập lớn nhất của các chủ nhà xe:")
    format_run(r, font_name="Segoe UI", size_pt=10.5)

    add_heading_2(doc, "5.1. Vấn đề 1: Mất kiểm soát trạng thái phương tiện theo thời gian thực")
    add_bullet(doc, "Thực trạng và Hậu quả", "Quản lý bằng sổ hoặc trí nhớ khiến nhân viên không biết chính xác xe nào đang rảnh tại bãi, xe nào đang cho thuê, xe nào sắp trả hoặc đang bảo dưỡng. Dẫn đến trùng lịch cho thuê hoặc từ chối khách dù bãi vẫn còn xe trống.")
    add_bullet(doc, "Giải pháp DriveRent", "Sơ đồ trạng thái xe trực quan với 3 màu nhận diện chuẩn gồm Sẵn sàng (xanh lá), Đang thuê (xanh dương), Đang bảo dưỡng (cam) kèm bộ lọc kiểm tra xe trống theo ngày giờ nhận và trả tức thì.")

    add_heading_2(doc, "5.2. Vấn đề 2: Tranh chấp hư hỏng trầy xước và Thất lạc giấy tờ tùy thân")
    add_bullet(doc, "Thực trạng và Hậu quả", "Lúc giao xe không ghi nhận kỹ vết trầy xước cũ; lúc khách trả xe phát sinh va quẹt mới thì hai bên cãi vã vì thiếu bằng chứng đối chiếu. Giữ giấy tờ gốc dễ thất lạc, khi có phạt nguội vi phạm giao thông không tra cứu lại được người lái.")
    add_bullet(doc, "Giải pháp DriveRent", "Biên bản bàn giao và thu hồi xe số hóa với danh mục kiểm tra 6 hạng mục, lưu trữ ảnh chụp 4 góc xe, ghi nhận số km đồng hồ, vạch xăng từ 1 đến 8 vạch và lưu ảnh chụp căn cước công dân, bằng lái xe gốc gắn liền với từng mã hợp đồng.")

    add_heading_2(doc, "5.3. Vấn đề 3: Tính toán chi phí thủ công và Nhập nhằng tiền cọc")
    add_bullet(doc, "Thực trạng và Hậu quả", "Khách thuê linh hoạt như trả trễ giờ, hụt xăng, thuê thêm mũ, phí giao xe tận nơi. Tính nhẩm thủ công dễ sai sót. Kế toán nhầm lẫn tiền cọc vào doanh thu dẫn đến bất đồng gay gắt khi thanh lý hợp đồng.")
    add_bullet(doc, "Giải pháp DriveRent", "Hệ thống tự động tính tiền và phân định tiền cọc. Tiền cọc lưu ở quỹ riêng, hệ thống tự động bù trừ vào tiền thuê thực tế hoặc hoàn cọc toàn bộ qua chuyển khoản nhanh.")

    add_heading_2(doc, "5.4. Vấn đề 4: Bỏ quên lịch thay nhớt và bảo dưỡng định kỳ")
    add_bullet(doc, "Thực trạng và Hậu quả", "Xe cho thuê chạy cường độ cao liên tục, không ai theo dõi số km kể từ lần thay nhớt trước. Hậu quả là động cơ xuống cấp nhanh, bó máy, mòn phanh gây nguy hiểm cho du khách trên đường đèo dốc và tốn hàng chục triệu đại tu.")
    add_bullet(doc, "Giải pháp DriveRent", "Cơ chế cảnh báo bảo trì kép tự động đối chiếu số km sau mỗi chuyến đi (xe máy mỗi 3.000 km, ô tô mỗi 5.000 km) và mốc ngày định kỳ để phát cảnh báo kịp thời.")

    # ==========================
    # CÂU 6: NHỮNG YÊU CẦU CHÍNH LÀ GÌ?
    # ==========================
    add_heading_1(doc, "6. Những yêu cầu chính là gì?")
    
    add_heading_2(doc, "6.1. Yêu cầu Phân hệ Khách hàng")
    add_bullet(doc, "Ngôn ngữ thiết kế kính mờ trong suốt toàn diện", "Toàn bộ thẻ xe, thanh điều hướng, cửa sổ thông báo, bảng thông số kỹ thuật đều sử dụng hiệu ứng kính mờ trong suốt, viền phản chiếu ánh sáng và đổ bóng cao cấp.")
    add_bullet(doc, "Thanh điều hướng cố định thông minh", "Cố định ở mép trên màn hình khi cuộn trang. Hiển thị dạng kính trong suốt khi ở đầu trang và tự động chuyển sang màu trắng thanh lịch khi cuộn xuống.")
    add_bullet(doc, "Phông nền phòng trưng bày gara vòm đèn LED", "Hình nền gara tràn toàn màn hình. Hiệu ứng xe tự động lăn bánh vào tâm bãi đỗ thay thế xe trước đó, thông tin xe hiển thị 2 bên và bên dưới trên các tấm kính mờ.")
    add_bullet(doc, "Tìm kiếm, Lọc xe và Đặt giữ chỗ 10 phút", "Lọc theo loại xe ô tô hoặc xe máy, hãng, số chỗ, giá thuê; cửa sổ thanh toán quét mã QR giữ xe trong 10 phút với đồng hồ đếm ngược thời gian thực.")
    add_bullet(doc, "Tra cứu đơn thuê trực tuyến", "Du khách chỉ cần nhập số điện thoại hoặc mã đặt chỗ để kiểm tra trạng thái duyệt xe và lịch nhận xe.")

    add_heading_2(doc, "6.2. Yêu cầu Phân hệ Quản trị Vận hành")
    add_bullet(doc, "Sơ đồ trạng thái đội xe", "Hỗ trợ 2 chế độ hiển thị thẻ lưới và bảng tổng hợp, 3 màu trạng thái chuẩn, hiển thị mức xăng từ 1 đến 8 vạch, số km đồng hồ, biển số xe.")
    add_bullet(doc, "Quy trình lập Hợp đồng 4 bước", "Bước 1: Chọn xe và thời gian, tự động kiểm tra trùng lịch -> Bước 2: Thông tin khách hàng -> Bước 3: Lưu ảnh căn cước công dân và bằng lái xe -> Bước 4: Tiền cọc và Phụ phí.")
    add_bullet(doc, "Biên bản bàn giao và Thu hồi xe số hóa", "Danh mục kiểm tra 6 hạng mục gồm thân vỏ, lốp, đèn còi, gương, xăng, phụ kiện; chụp ảnh hiện trạng 4 góc; nhập phí phạt đền bù hư hại.")
    add_bullet(doc, "Quyết toán thanh lý hợp đồng", "Tự động tính tiền theo giờ hoặc ngày thực tế, phân tách tiền cọc, tự động tính số tiền hoàn trả hoặc thu thêm.")
    add_bullet(doc, "Quản lý bảo dưỡng và Nhắc thay nhớt", "Theo dõi số km lăn bánh, khóa xe không cho thuê khi đang bảo dưỡng, tự động sinh mốc km bảo dưỡng tiếp theo.")
    add_bullet(doc, "Mẫu in hợp đồng khổ A4", "Mẫu in pháp lý chuẩn gồm quốc hiệu, trách nhiệm bên thuê và cho thuê, chữ ký xác nhận 2 bên.")

    # ==========================
    # CÂU 7: PHẠM VI CÔNG VIỆC ĐẾN ĐÂU?
    # ==========================
    add_heading_1(doc, "7. Phạm vi công việc đến đâu?")
    
    add_heading_2(doc, "7.1. Trong phạm vi thực hiện")
    add_bullet(doc, "Ứng dụng giao diện web", "Xây dựng toàn bộ giao diện tương thích hoàn hảo từ điện thoại, máy tính bảng đến máy tính để bàn và màn hình lớn.")
    add_bullet(doc, "Hệ thống thiết kế kính mờ trong suốt", "Áp dụng phong cách kính mờ đồng bộ cho 7 trang khách hàng và 10 màn hình quản trị nội bộ; chuẩn hóa hình ảnh xe cắt nền trong suốt và hình nền gara đèn LED.")
    add_bullet(doc, "Nghiệp vụ cốt lõi và Quản lý trạng thái", "Lập trình đầy đủ luồng dữ liệu thời gian thực: Đặt xe trực tuyến -> Duyệt đơn -> Lập hợp đồng -> Giao xe -> Nhận xe trả -> Quyết toán cọc -> Bảo dưỡng định kỳ.")
    add_bullet(doc, "Dữ liệu mẫu và Mẫu in hợp đồng A4", "Cung cấp sẵn dữ liệu thực tế gồm 12 đầu xe ô tô và xe máy, khách hàng, hợp đồng, lịch bảo trì và mẫu in hợp đồng pháp lý A4.")
    add_bullet(doc, "Đóng gói và Kiểm thử", "Kiểm thử lệnh đóng gói không lỗi và đẩy toàn bộ mã nguồn lên GitHub.")

    add_heading_2(doc, "7.2. Ngoài phạm vi hiện tại")
    add_bullet(doc, "Máy chủ và Cơ sở dữ liệu đám mây", "Kết nối hệ thống cơ sở dữ liệu phân tán lưu trữ lâu dài trên đám mây.")
    add_bullet(doc, "Cổng thanh toán tự động", "Tích hợp kết nối ngân hàng để tự động xác nhận tiền cọc mà không cần nhân viên kiểm tra thủ công.")
    add_bullet(doc, "Định vị phần cứng hành trình", "Lắp đặt thiết bị định vị trên xe để tự động đọc số km đồng hồ và vị trí xe theo thời gian thực.")
    add_bullet(doc, "Nhận dạng tự động giấy tờ tùy thân", "Tự động đọc và trích xuất thông tin họ tên, số căn cước từ ảnh chụp giấy tờ.")

    # ==========================
    # CÂU 8: KẾT QUẢ CẦN ĐẠT ĐƯỢC LÀ GÌ?
    # ==========================
    add_heading_1(doc, "8. Kết quả cần đạt được là gì?")
    
    add_heading_2(doc, "8.1. Sản phẩm bàn giao cụ thể")
    add_bullet(doc, "1. Kho mã nguồn hoàn chỉnh trên GitHub", "Được tổ chức sạch sẽ, cấu trúc rõ ràng, đóng gói không có lỗi.")
    add_bullet(doc, "2. Ứng dụng web hoàn chỉnh", "Chạy mượt mà cả Cổng Khách hàng và Cổng Quản trị với thời gian phản hồi tức thì dưới 100 phần nghìn giây.")
    add_bullet(doc, "3. Bộ tài liệu đặc tả chuẩn mực", "Tệp văn bản tóm tắt dự án định dạng chuyên nghiệp theo chuẩn 8 câu hỏi cốt lõi.")

    add_heading_2(doc, "8.2. Chỉ số hiệu quả kỳ vọng")
    add_bullet(doc, "Hiệu suất vận hành", "Giảm 85% thời gian làm thủ tục lập hợp đồng và bàn giao nhận xe tại quầy, từ 15–20 phút xuống dưới 2–3 phút.")
    add_bullet(doc, "Chính xác tài chính", "Đạt 100% độ chính xác trong việc tính cước thuê và quyết toán hoàn cọc; loại trừ hoàn toàn tình trạng nhầm lẫn tiền cọc vào doanh thu.")
    add_bullet(doc, "Bảo vệ tài sản phương tiện", "100% hợp đồng có lưu ảnh đối chiếu 4 góc xe và ảnh giấy tờ tùy thân; 100% đầu xe được nhắc thay nhớt đúng hạn.")
    add_bullet(doc, "Trải nghiệm và Chuyển đổi", "Giao diện kính mờ trong suốt chuẩn showroom tạo ấn tượng chuyên nghiệp vượt trội, nâng cao tỷ lệ đặt giữ xe của khách du lịch.")

    # Save document
    doc.save(filename)
    print(f"Successfully generated clean Vietnamese Word document: {filename}")

if __name__ == "__main__":
    out_path = os.path.abspath("PROJECT_BRIEF.docx")
    create_brief_docx(out_path)
